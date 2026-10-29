"""Generate synthetic 6-well plate scans so the pipeline can be tried without any real data.

    python examples/make_synthetic_plates.py --out examples/synthetic [--seed 0]

Writes two plates (control: ~60 colonies per well; treated: ~30 colonies per well, smaller colonies),
300 dpi JPEG with DPI metadata, plus a filled PlateMap.xlsx. Colonies are drawn as soft purple discs
with log-normal sizes on a lavender well background with a bright rim ring, mimicking a crystal violet
stained plate on a flatbed scanner. Intended for testing the software, not for any biological inference.
"""
import argparse, os, numpy as np
from PIL import Image
DPI = 300; PX_MM = DPI / 25.4; PITCH = 39.12; WELL_R = 17.4

def plate(n_per_well, diam_mm, seed):
    rng = np.random.default_rng(seed); W, H = 2000, 1500
    img = np.full((H, W, 3), (235, 233, 236), np.float32)            # scanner bed
    yy, xx = np.mgrid[:H, :W].astype(np.float32); cx0, cy0 = W / 2, H / 2
    centres = [(cx0 + (c - 1) * PITCH * PX_MM, cy0 + (r - 0.5) * PITCH * PX_MM) for r in range(2) for c in range(3)]
    R = WELL_R * PX_MM; counts = []
    for cx, cy in centres:
        d = np.hypot(xx - cx, yy - cy)
        img[d <= R] = (208, 196, 232)                                   # lavender well floor
        ring = (d > R - 2) & (d <= R + 3); img[ring] = (250, 248, 252)   # bright rim ring
        wall = (d > R + 3) & (d <= R + 12); img[wall] = (215, 212, 220)
        n = rng.poisson(n_per_well); counts.append(n)
        for _ in range(n):
            rr = 0.93 * R * np.sqrt(rng.uniform(0, 1)); th = rng.uniform(0, 2 * np.pi); x, y = cx + rr * np.cos(th), cy + rr * np.sin(th)
            dia = rng.lognormal(np.log(diam_mm), 0.35) * PX_MM; s = dia / 2
            m = np.exp(-((xx - x) ** 2 + (yy - y) ** 2) / (2 * (s / 1.6) ** 2)); m = np.clip(m * 1.4, 0, 1)
            dark = np.array((95, 60, 165), np.float32); img = img * (1 - m[..., None] * 0.95) + dark * m[..., None] * 0.95
    img += rng.normal(0, 2.5, img.shape); img = np.clip(img, 0, 255).astype(np.uint8)
    return Image.fromarray(img), counts

if __name__ == "__main__":
    ap = argparse.ArgumentParser(); ap.add_argument("--out", default="examples/synthetic"); ap.add_argument("--seed", type=int, default=0); a = ap.parse_args()
    os.makedirs(a.out, exist_ok=True); truth = []
    for name, n, dia, s in [("control", 60, 0.9, a.seed), ("treated", 30, 0.7, a.seed + 1)]:
        im, counts = plate(n, dia, s); im.save(f"{a.out}/synthetic_{name}.jpg", quality=92, dpi=(DPI, DPI)); truth.append((name, counts))
        print(f"synthetic_{name}.jpg: colonies drawn per well {counts}")
    import sys; sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "colonycounter"))
    from cfa_platemap import make_template
    from openpyxl import load_workbook
    make_template(a.out, f"{a.out}/PlateMap.xlsx", 6)
    wb = load_workbook(f"{a.out}/PlateMap.xlsx"); ws = wb["PlateMap"]
    for row in ws.iter_rows():
        if row[0].value == "File":
            cond = "control" if "control" in row[1].value else "treated"; r = row[0].row; ws.cell(r + 1, 2, "SynthLine")
            for i in range(2):
                for c in range(3): ws.cell(r + 6 + i, 2 + c, cond)
    wb["Controls"].append(["SynthLine", "control"]); wb.save(f"{a.out}/PlateMap.xlsx")
    with open(f"{a.out}/ground_truth.txt", "w") as f:
        for name, counts in truth: f.write(f"{name}\t" + "\t".join(map(str, counts)) + "\n")
    print("plate map written:", f"{a.out}/PlateMap.xlsx")
